"""Offline verification for every bundled resume sample."""

from __future__ import annotations

import argparse
import json
import os
from pathlib import Path
import re
import sys
import tempfile

# Never send bundled candidate documents to an external provider during verification.
os.environ["LLM_PROVIDER"] = "off"

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from backend.app.resume_pipeline import ResumePipeline
from backend.app.schemas import CandidateProfile
RESUME_DIRS = (ROOT / "data" / "resumes" / "PDF", ROOT / "data" / "resumes" / "Word")
OUTPUT_DIR = ROOT / "output" / "resumes"


def create_sample_jd(path: Path) -> None:
    from docx import Document

    document = Document()
    document.add_heading("Role Overview", level=1)
    document.add_paragraph("The selected candidate will design software and work with customers.")
    document.add_heading("Key Responsibilities", level=1)
    document.add_paragraph("Collaborate with the product team and document project decisions.")
    document.save(path)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--write-output",
        action="store_true",
        help="Write parsed profiles to the ignored output/resumes folder.",
    )
    args = parser.parse_args()

    files = sorted(
        path
        for directory in RESUME_DIRS
        if directory.is_dir()
        for path in directory.iterdir()
        if path.is_file() and path.suffix.lower() in {".pdf", ".docx"}
    )
    if not files:
        print("No bundled PDF or DOCX resumes were found.", file=sys.stderr)
        return 2

    pipeline = ResumePipeline()
    failures: list[str] = []
    for source in files:
        try:
            profile = pipeline.process_resume(str(source), source_filename=source.name)
            validated = CandidateProfile.model_validate(profile)
            if validated.source_file != source.name:
                raise ValueError("The uploaded source filename was not preserved.")
            if args.write_output:
                OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
                output_name = f"{re.sub(r'[^a-z0-9]+', '_', source.stem.lower()).strip('_')}_{source.suffix[1:].lower()}.json"
                (OUTPUT_DIR / output_name).write_text(
                    json.dumps(validated.model_dump(mode="json"), ensure_ascii=False, indent=2) + "\n",
                    encoding="utf-8",
                )
            print(f"PASS {source.relative_to(ROOT)}: {len(validated.skills)} extracted skills")
        except Exception as exc:
            failures.append(source.name)
            print(f"FAIL {source.name}: {exc}", file=sys.stderr)

    # Verify that an ordinary JD is rejected by the resume-specific endpoint.
    try:
        from backend.app.document_parser import ResumeDocumentParser

        with tempfile.TemporaryDirectory(prefix="radix-jd-guardrail-", dir=ROOT / "output") as temp_dir:
            jd_path = Path(temp_dir) / "sample-jd.docx"
            create_sample_jd(jd_path)
            try:
                ResumeDocumentParser().parse(str(jd_path))
                failures.append("resume/JD mismatch guardrail")
                print("FAIL resume/JD mismatch guardrail accepted a JD.", file=sys.stderr)
            except ValueError as exc:
                if "Document mismatch guardrail" not in str(exc):
                    raise
                print("PASS resume/JD mismatch guardrail")
    except Exception as exc:
        failures.append("resume/JD mismatch guardrail")
        print(f"FAIL resume/JD mismatch guardrail: {exc}", file=sys.stderr)

    if failures:
        print(f"Verification failed: {len(failures)} issue(s).", file=sys.stderr)
        return 1
    print(f"Verification passed: {len(files)} resumes plus the JD upload guardrail.")
    if args.write_output:
        print(f"Profiles written to {OUTPUT_DIR}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
